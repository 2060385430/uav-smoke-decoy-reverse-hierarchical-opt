import os
import csv
import json
import pickle
import importlib.util
from datetime import datetime

import numpy as np
import openpyxl
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "ablation_rho.py")
ABL_JSON = os.path.join(HERE, "ablation_results.json")
PKL = os.path.join(HERE, "..", "candidate_lib.pkl")
CSV_CHAMP = os.path.join(HERE, "best_plan_rho.csv")
OUT_ANCHOR = os.path.join(HERE, "plan_anchor_20.46.xlsx")
OUT_CHAMP = os.path.join(HERE, "plan_champion_22.717.xlsx")

ANCHOR_ASSIGN = ("M1", "M2", "M1", "M1", "M3")
CHAMPION_KEY = "M1,M2,M1,M1,M3"

COLUMNS = ["无人机", "导弹", "速度", "航向角°", "投放时刻", "起爆延迟",
           "起爆时刻", "起爆X", "起爆Y", "起爆Z", "遮蔽区间", "时长"]

# ---- 导入 ablation_rho.py（不执行其 __main__） ----
spec = importlib.util.spec_from_file_location("ablation_rho", SRC)
rho = importlib.util.module_from_spec(spec)
spec.loader.exec_module(rho)


def detonation_state(drone, theta, v, drop, det):
    """投放/起爆运动学（与 ablation_rho.shielded_scalar 同一套公式）"""
    init = rho.DRONES[drone]
    uav = np.array([np.cos(theta), np.sin(theta), 0.0])
    drop_pos = init + v * drop * uav
    det_xy = drop_pos[:2] + v * det * uav[:2]
    det_z = drop_pos[2] - 0.5 * rho.G * det * det
    return det_xy[0], det_xy[1], det_z


def build_anchor_rows(lib):
    """锚点方案：各机第 0 候选，逐弹区间 fine 精度重算"""
    rows = []
    per_m = {m: [] for m in rho.MISSILE_LIST}
    for i, drone in enumerate(rho.DRONE_LIST):
        m = ANCHOR_ASSIGN[i]
        c = lib[(drone, m)][0]  # 第 0 候选
        v, theta = float(c["v"]), float(c["theta"])
        for drop, det in zip(c["drop_times"], c["det_delays"]):
            drop, det = float(drop), float(det)
            dx, dy, dz = detonation_state(drone, theta, v, drop, det)
            iv = rho.interval_for(drone, m, theta, v, drop, det,
                                  rho.E_FINE, 0.005, bisect=40)
            if iv is not None:
                per_m[m].append(iv)
                iv_str = f"[{iv[0]:.4f},{iv[1]:.4f}]"
                dur = round(iv[1] - iv[0], 4)
            else:
                iv_str, dur = "无", 0.0
            rows.append([drone, m, round(v, 4),
                         round(float(np.degrees(theta)), 4),
                         round(drop, 4), round(det, 4), round(drop + det, 4),
                         round(dx, 2), round(dy, 2), round(dz, 2),
                         iv_str, dur])
    per_m_dur = {m: rho.merge(ivs) for m, ivs in per_m.items()}
    return rows, per_m_dur


def read_champion_rows():
    """冠军方案：直接读取 best_plan_rho.csv 逐弹数据"""
    rows = []
    with open(CSV_CHAMP, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        header = next(reader)
        assert header == COLUMNS, f"csv 列名变化: {header}"
        for r in reader:
            row = []
            for j, cell in enumerate(r):
                if j in (0, 1, 10):        # 无人机/导弹/遮蔽区间：文本
                    row.append(cell)
                else:                       # 其余：数值
                    row.append(float(cell))
            rows.append(row)
    return rows


def write_xlsx(path, rows, sheet_name, summary_rows, note):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    widths = [8, 6, 10, 10, 10, 10, 10, 11, 10, 10, 20, 9]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w

    ws2 = wb.create_sheet("汇总")
    ws2.append(["项目", "时长(s)"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for name, val in summary_rows:
        ws2.append([name, val])
    ws2.append([])
    ws2.append(["说明", note])
    ws2.append(["生成时间", datetime.now().isoformat(timespec="seconds")])
    ws2.column_dimensions["A"].width = 34
    ws2.column_dimensions["B"].width = 70
    wb.save(path)
    print(f"[落盘] {path}", flush=True)


def main():
    with open(ABL_JSON, "r", encoding="utf-8") as f:
        st = json.load(f)
    anchor_fine_stored = st["anchor"]["fine_total"]           # 20.459533
    champ_fine_stored = st["refines"][CHAMPION_KEY]["fine_total"]  # 22.717447

    # ---------- 1) 锚点方案 20.46 s ----------
    with open(PKL, "rb") as f:
        lib = pickle.load(f)
    rows_a, per_m_a = build_anchor_rows(lib)
    total_a = sum(per_m_a.values())
    summary_a = ([(f"{m} 遮蔽并集时长", round(per_m_a[m], 6))
                  for m in rho.MISSILE_LIST]
                 + [("三导弹合计（本表重算）", round(total_a, 6)),
                    ("ablation_results.json 存档 anchor.fine_total",
                     round(anchor_fine_stored, 6)),
                    ("重算与存档一致", bool(abs(total_a - anchor_fine_stored) < 1e-6))])
    note_a = ("候选库锚点方案：分配 M1,M2,M1,M1,M3、各机第 0 候选（= 消融 A1 未精修方案）；"
              "逐弹参数取自 candidate_lib.pkl，遮蔽区间按 fine 精度"
              "（200点/圆周、dt=0.005、二分40次）用 ablation_rho.interval_for 重算。")
    write_xlsx(OUT_ANCHOR, rows_a, "锚点方案20.46s逐弹明细", summary_a, note_a)
    print(f"  锚点方案：三导弹合计 {total_a:.6f} s"
          f"（存档 {anchor_fine_stored:.6f}，一致: {abs(total_a - anchor_fine_stored) < 1e-6}）",
          flush=True)

    # ---------- 2) 冠军方案 22.717 s ----------
    rows_c = read_champion_rows()
    per_m_c = {m: 0.0 for m in rho.MISSILE_LIST}
    for r in rows_c:
        per_m_c[r[1]] += r[11]  # 各弹时长按导弹归类（csv 中为 fine 区间时长）
    summary_c = ([(f"{m} 各弹时长合计（csv 逐弹求和）", round(per_m_c[m], 6))
                  for m in rho.MISSILE_LIST]
                 + [("三导弹合计（csv 逐弹求和）", round(sum(per_m_c.values()), 6)),
                    ("ablation_results.json 存档 fine_total",
                     round(champ_fine_stored, 6)),
                    ("求和与存档一致(±1e-3 舍入内)",
                     bool(abs(sum(per_m_c.values()) - champ_fine_stored) < 1e-3))])
    note_c = ("RHO 冠军方案：ablation_results.json 中 refines[\"M1,M2,M1,M1,M3\"].x；"
              "逐弹数据由 best_plan_rho.csv 原样转写（fine 精度明细）。")
    write_xlsx(OUT_CHAMP, rows_c, "冠军方案22.717s逐弹明细", summary_c, note_c)
    print(f"  冠军方案：csv 逐弹求和 {sum(per_m_c.values()):.4f} s"
          f"（存档 {champ_fine_stored:.6f}）", flush=True)


if __name__ == "__main__":
    main()
